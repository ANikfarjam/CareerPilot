"""
Turns a resume draft's Markdown into a compact, one-page Word document, styled like a classic
single-column resume. drive.py uploads it to Google Drive, which converts it to a Google Doc and
keeps the styling (margins, fonts, right-aligned dates, section rules, bullets, links).

Reads the layout RESUME_PROMPT asks for (agents/resume_builder/agent.py):
    # Name                          -> centered, bold, large
    contact | line | ...            -> centered (first line after the name)
    ## Section                      -> bold caps with a rule under it
    ### Title — Employer | ... | Dates -> bold title, dates right-aligned at the margin
    *Location | Dates*              -> (right under a ###) joined into that entry's line
    - bullet                        -> tight bullet
    anything else                   -> plain paragraph
Inline **bold**, *italic*, `code`, [text](url), <url> and bare URLs are kept. Drafts the user
edited out of that layout still render, as plain paragraphs.

One page: the body font steps down (10 -> 9.5 -> 9pt) until an estimate of the rendered height
fits the page. It's an estimate, so a very long draft can still spill onto a second page.

Run it:
    python -m agents.resume_builder.resume_docx resume.md --out resume.docx
"""

import argparse
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

FONT = "Arial"
MARGIN = Inches(0.5)
TEXT_WIDTH = Inches(7.5)  # Letter width minus margins
BULLET_INDENT = Inches(0.25)
BODY_SIZES = (10, 9.5, 9)  # tried in order until the estimate fits one page
NAME_SCALE, SECTION_SCALE = 1.6, 1.05  # relative to the body size
SECTION_SPACE_BEFORE, SECTION_SPACE_AFTER = 6, 2  # pt
PAGE_HEIGHT_PT = 10 * 72  # 11in minus margins
LINE_HEIGHT = 1.15  # Arial single spacing, in font sizes
CHARS_PER_LINE_AT_10PT = 112  # Arial 10pt across TEXT_WIDTH
LINK_COLOR = RGBColor(0x11, 0x55, 0xCC)
RULE_COLOR = "808080"

DATE_RE = re.compile(r"\b(19|20)\d{2}\b|\b(present|current|expected)\b", re.I)
INLINE_RE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|\[(?P<text>[^\]]+)\]\((?P<href>[^)\s]+)\)"
    r"|<(?P<auto>https?://[^>\s]+)>"
    r"|`(?P<code>[^`]+)`"
    r"|(?<![\w*])\*(?!\s)(?P<em>.+?)(?<!\s)\*(?![\w*])"
    r"|(?<!\w)_(?!\s)(?P<em2>.+?)(?<!\s)_(?!\w)"
    r"|(?P<url>https?://[^\s,|)]+)"
)
TITLE_SPLIT_RE = re.compile(r"\s+[—–]\s+")

# pPr children that must come after w:pBdr (OOXML schema order).
_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
               "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
               "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
               "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
               "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


# ---------- parsing ----------

def parse(markdown: str) -> list[tuple[str, str]]:
    """The draft as (kind, text) blocks; kind is name / contact / section / entry / bullet / text."""
    blocks: list[tuple[str, str]] = []
    expect_contact = False
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            continue
        if line.startswith("# "):
            blocks.append(("name", line[2:].strip()))
            expect_contact = True
            continue
        if line.startswith("## "):
            blocks.append(("section", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("entry", line[4:].strip()))
        elif line.startswith("|"):
            # Markdown table: each row becomes a bullet; the separator row drops the header above it
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                if blocks and blocks[-1][0] == "row":
                    blocks.pop()
            else:
                blocks.append(("row", " – ".join(c for c in cells if c)))
        elif re.match(r"[-+•●]\s+|\*\s+", line):
            blocks.append(("bullet", re.sub(r"^[-*+•●]\s+", "", line)))
        elif blocks and blocks[-1][0] == "entry" and (meta := _entry_meta(line)):
            # "*Location | Dates*" on the line under a ### heading: join it so the date right-aligns
            blocks[-1] = ("entry", f"{blocks[-1][1]} | {meta}")
        elif expect_contact:
            blocks.append(("contact", line))
        else:
            blocks.append(("text", line))
        expect_contact = False
    return [("bullet" if kind == "row" else kind, text) for kind, text in blocks]


def _plain(text: str) -> str:
    """The text as it renders, without Markdown markers (for the length estimate)."""
    return INLINE_RE.sub(lambda m: m["bold"] or m["text"] or m["em"] or m["em2"] or m["code"]
                         or _link_label(m["auto"] or m["url"]), text)


def _link_label(url: str) -> str:
    """A bare link as shown: without the scheme and trailing slash, like a printed resume."""
    return re.sub(r"^https?://(www\.)?", "", url).rstrip("/")


def _entry_meta(line: str) -> str:
    """The line's text if it is an entry's dated detail line ('*Stanford, CA | Sep 2025 – Present*')."""
    inner = re.fullmatch(r"([*_])(.+)\1", line)
    text = inner[2].strip() if inner else line
    return text if DATE_RE.search(text) and len(text) < 80 else ""


def _split_entry(text: str) -> tuple[str, str, str]:
    """'Title — Employer | Location | Dates' -> (bold title, rest with its separator, date)."""
    parts = [p.strip() for p in text.split("|") if p.strip()]
    date = parts.pop() if len(parts) > 1 and DATE_RE.search(parts[-1]) and len(parts[-1]) < 40 else ""
    head, tail = (parts[0] if parts else ""), parts[1:]
    rest = f", {', '.join(tail)}" if tail else ""
    title = TITLE_SPLIT_RE.split(head, maxsplit=1)
    if len(title) == 2:
        head, rest = title[0], f" – {title[1]}{rest}"
    return head, rest, date


# ---------- fitting ----------

def estimate_height(blocks: list[tuple[str, str]], size: float) -> float:
    """Rough rendered height in points at this body size."""
    per_line = CHARS_PER_LINE_AT_10PT * 10 / size
    total = 0.0
    for kind, text in blocks:
        plain = _plain(text)
        if kind == "name":
            total += size * NAME_SCALE * LINE_HEIGHT
        elif kind == "section":
            total += size * SECTION_SCALE * LINE_HEIGHT + SECTION_SPACE_BEFORE + SECTION_SPACE_AFTER
        else:
            width = per_line * (7.25 / 7.5 if kind == "bullet" else 1)
            lines = max(1, -(-len(plain) // int(width)))  # ceil
            total += lines * size * LINE_HEIGHT
    return total


def fit_size(blocks: list[tuple[str, str]]) -> float:
    """Largest body size whose estimate fits one page (the smallest one if none does)."""
    for size in BODY_SIZES:
        if estimate_height(blocks, size) <= PAGE_HEIGHT_PT:
            return size
    return BODY_SIZES[-1]


# ---------- rendering ----------

def _hyperlink(par, text: str, url: str, size: float, bold: bool = False) -> None:
    run = par.add_run(text)
    run.bold = bold or None
    run.font.size = Pt(size)
    run.font.color.rgb = LINK_COLOR
    run.font.underline = True
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    link.append(run._r)  # moves the run inside the link
    par._p.append(link)


def _inline(par, text: str, size: float, bold: bool = False) -> None:
    """Adds text to the paragraph, turning its inline Markdown into runs and links."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], size, bold)
        if m["bold"]:
            _run(par, m["bold"], size, True)
        elif m["code"]:
            _run(par, m["code"], size, bold)
        elif m["href"]:
            _hyperlink(par, m["text"], m["href"], size, bold)
        elif m["auto"] or m["url"]:
            url = m["auto"] or m["url"]
            _hyperlink(par, _link_label(url), url, size, bold)
        else:
            _run(par, m["em"] or m["em2"], size, bold, italic=True)
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], size, bold)


def _run(par, text: str, size: float, bold: bool = False, italic: bool = False) -> None:
    run = par.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold or None
    run.italic = italic or None


def _bottom_rule(par) -> None:
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in {"val": "single", "sz": "4", "space": "1", "color": RULE_COLOR}.items():
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    par._p.get_or_add_pPr().insert_element_before(border, *_AFTER_PBDR)


def _setup(doc, size: float) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = MARGIN
    for name in ("Normal", "List Bullet"):
        style = doc.styles[name]
        style.font.name = FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        fmt = style.paragraph_format
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0


def build(blocks: list[tuple[str, str]], size: float) -> Document:
    doc = Document()
    _setup(doc, size)
    for kind, text in blocks:
        if kind == "bullet":
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.left_indent = BULLET_INDENT
            par.paragraph_format.first_line_indent = -Inches(0.17)
            _inline(par, text, size)
            continue
        par = doc.add_paragraph()
        fmt = par.paragraph_format
        if kind == "name":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _inline(par, text, size * NAME_SCALE, bold=True)
        elif kind == "contact":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _inline(par, text, size)
        elif kind == "section":
            fmt.space_before, fmt.space_after = Pt(SECTION_SPACE_BEFORE), Pt(SECTION_SPACE_AFTER)
            fmt.keep_with_next = True
            _bottom_rule(par)
            _inline(par, _plain(text).upper(), size * SECTION_SCALE, bold=True)
        elif kind == "entry":
            head, rest, date = _split_entry(text)
            fmt.keep_with_next = True
            _inline(par, head, size, bold=True)
            if rest:
                _inline(par, rest, size)
            if date:
                fmt.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
                _inline(par, f"\t{date}", size)
        else:
            _inline(par, text, size)
    return doc


def markdown_to_docx(markdown: str) -> bytes:
    """The resume Markdown as a styled .docx (bytes), sized to fit one page."""
    blocks = parse(markdown)
    out = io.BytesIO()
    build(blocks, fit_size(blocks)).save(out)
    return out.getvalue()


def main() -> None:
    parser = argparse.ArgumentParser(description="Render a Markdown resume as a styled one-page .docx.")
    parser.add_argument("markdown", type=Path, help="resume Markdown file")
    parser.add_argument("--out", type=Path, required=True, help="where to write the .docx")
    args = parser.parse_args()
    blocks = parse(args.markdown.read_text())
    size = fit_size(blocks)
    args.out.write_bytes(markdown_to_docx(args.markdown.read_text()))
    print(f"{args.out} ({size}pt body, ~{estimate_height(blocks, size) / PAGE_HEIGHT_PT:.0%} of a page)")


if __name__ == "__main__":
    main()
